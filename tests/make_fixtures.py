"""Create test fixtures: a .docx (python-docx) and a .md sermon."""
import docx, pathlib
here = pathlib.Path(__file__).parent / 'fixtures'
d = docx.Document()
d.add_heading('The Good Shepherd', level=1)
d.add_heading('Point 1: He knows your name', level=2)
p = d.add_paragraph('Turn with me to John 10:11. ')
r = p.add_run('This is bold.'); r.bold = True
p.add_run(' And ')
r = p.add_run('this is italic.'); r.italic = True
d.add_paragraph('First list item referencing Psalm 23', style='List Bullet')
d.add_paragraph('Second list item referencing Isa 40:11', style='List Bullet')
d.add_paragraph('Numbered one', style='List Number')
d.add_paragraph('Numbered two', style='List Number')
d.save(here / 'test-sermon.docx')
(here / 'test-sermon.md').write_text('# Markdown Sermon\n\n**Bold** and *italic* with Rom 12:1-2.\n\n- one\n- two\n\n1. first\n2. second\n')
print('fixtures ok')

# Highlighted .docx for big-screen slide extraction.
from docx.enum.text import WD_COLOR_INDEX
h = docx.Document()
h.add_heading('Anchored', level=1)
p = h.add_paragraph('Life will test you. ')
r = p.add_run('Hope is only as strong as what it is anchored to.'); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
p.add_run(' That is the point.')
p = h.add_paragraph('Normal text with no highlight at all.')
p = h.add_paragraph('Split runs: ')
r = p.add_run('Jesus is '); r.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
r = p.add_run('the anchor'); r.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN; r.bold = True
p.add_run(' and more plain text.')
h.add_heading('Point 2: Hold on', level=2)
h.save(here / 'highlight-sermon.docx')
print('highlight fixture ok')
